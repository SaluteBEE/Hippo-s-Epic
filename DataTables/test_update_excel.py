"""
update_excel.py 的完整单元测试
验证导出功能是否生成了正确的文件和数据到 Excel
运行: python -m pytest DataTables/test_update_excel.py -v
  或: python -m unittest DataTables.test_update_excel -v
"""

import unittest
import os
import sys
import json
import tempfile
import shutil
import codecs

sys.path.insert(0, os.path.dirname(__file__))
from update_excel import update_animationstate, update_slotstate


def _write_json(temp_dir, data, filename="test_data.json"):
    json_path = os.path.join(temp_dir, filename)
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    return json_path


def _write_json_bom(temp_dir, data, filename="bom_test.json"):
    json_path = os.path.join(temp_dir, filename)
    with open(json_path, "wb") as f:
        f.write(codecs.BOM_UTF8)
        f.write(json.dumps(data, ensure_ascii=False).encode("utf-8"))
    return json_path


def _read_xlsx(xlsx_path):
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    rows = []
    for r in range(1, ws.max_row + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            row.append(ws.cell(r, c).value)
        rows.append(row)
    return rows, ws


class TestAnimationStateFreshCreate(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "animationstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_creates_file_when_not_exists(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)
        self.assertTrue(os.path.exists(self.xlsx_path))

    def test_headers_correct(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(
            rows[0], ["##var", "id", "personid", "statename", "prefabtype", "slotstateid"]
        )
        self.assertEqual(
            rows[1], ["##type", "int", "int#ref=cfg.person.TbPerson", "string", "int", "int#ref=cfg.slotstate.TbSlotstate"]
        )
        self.assertEqual(rows[2], ["##", None, None, None, None, None])
        self.assertEqual(rows[3][0], "##comment")

    def test_single_row_data(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(rows[4], [None, 1000, 1, "idle", 1, 0])

    def test_multiple_rows_sequential_ids(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 1, "statename": "talk", "prefabtype": 1},
                {"personid": 1, "statename": "think1", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(len(rows), 7)
        self.assertEqual(rows[4][1], 1000)
        self.assertEqual(rows[5][1], 1001)
        self.assertEqual(rows[6][1], 1002)

    def test_first_column_none(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertIsNone(rows[4][0])


class TestAnimationStateUpdate(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "animationstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_update_existing_slotstateid(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1, "slotstateid": 0},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1, "slotstateid": 5},
            ],
            "update.json",
        )
        update_animationstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(rows[4][5], 5)

    def test_update_preserves_id(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 1, "statename": "talk", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 1, "statename": "talk", "prefabtype": 1},
            ],
            "update.json",
        )
        update_animationstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(rows[4][1], 1000)
        self.assertEqual(rows[5][1], 1001)

    def test_add_new_rows_to_existing(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 2, "statename": "idle", "prefabtype": 1},
                {"personid": 2, "statename": "talk", "prefabtype": 1},
            ],
            "add.json",
        )
        update_animationstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(len(rows), 7)
        self.assertEqual(rows[4][1], 1000)
        self.assertEqual(rows[5][1], 2000)
        self.assertEqual(rows[6][1], 2001)

    def test_different_personid_same_statename(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 2, "statename": "idle", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(len(rows), 6)
        self.assertEqual(rows[4][3], "idle")
        self.assertEqual(rows[5][3], "idle")

    def test_id_auto_increments_from_max(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 1, "statename": "talk", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 2, "statename": "angry", "prefabtype": 1},
            ],
            "new.json",
        )
        update_animationstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(rows[6][1], 2000)


class TestAnimationStateCleanup(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "animationstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_delete_stale_rows_for_same_person_prefab(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 1, "statename": "talk", "prefabtype": 1},
                {"personid": 1, "statename": "angry", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
            ],
            "cleanup.json",
        )
        update_animationstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        data_rows = [r for r in rows[4:] if r[1] is not None]
        self.assertEqual(len(data_rows), 1)
        self.assertEqual(data_rows[0][3], "idle")

    def test_no_delete_different_person(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 1, "statename": "talk", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 2, "statename": "idle", "prefabtype": 1},
            ],
            "other.json",
        )
        update_animationstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        data_rows = [r for r in rows[4:] if r[1] is not None]
        self.assertEqual(len(data_rows), 3)

    def test_no_delete_different_prefabtype(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 1, "statename": "talk", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 2},
            ],
            "battle.json",
        )
        update_animationstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        data_rows = [r for r in rows[4:] if r[1] is not None]
        self.assertEqual(len(data_rows), 3)

    def test_delete_only_scope_rows(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 1, "statename": "talk", "prefabtype": 1},
                {"personid": 1, "statename": "angry", "prefabtype": 2},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
            ],
            "cleanup.json",
        )
        update_animationstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        data_rows = [r for r in rows[4:] if r[1] is not None]
        self.assertEqual(len(data_rows), 2)
        statenames = [r[3] for r in data_rows]
        self.assertIn("idle", statenames)
        self.assertIn("angry", statenames)
        self.assertNotIn("talk", statenames)


class TestAnimationStateBOM(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "animationstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_utf8_bom_json_handled(self):
        json_path = _write_json_bom(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(rows[4][3], "idle")

    def test_no_bom_json_handled(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(rows[4][3], "idle")


class TestAnimationStateMigration(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "animationstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_old_format_with_compositionname_auto_migrates(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "animationstate"
        ws.append(["##var", "id", "personid", "statename", "prefabtype", "compositionname", "slotstateid"])
        ws.append(["##type", "int", "int#ref=cfg.person.TbPerson", "string", "int", "string", "int"])
        ws.append(["##", None, None, None, None, None, None])
        ws.append(["##comment", "ID", "角色", "状态名", "预制体类型", "组合名", "插槽状态ID"])
        ws.append([None, 1, 1, "idle", 1, "idle", 0])
        ws.append([None, 2, 2, "talk", 1, "talk", 0])
        wb.save(self.xlsx_path)
        wb.close()

        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 2, "statename": "talk", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertNotIn("compositionname", rows[0])
        self.assertEqual(rows[0], ["##var", "id", "personid", "statename", "prefabtype", "slotstateid"])
        data_rows = [r for r in rows[4:] if r[1] is not None]
        self.assertEqual(len(data_rows), 2)


class TestSlotStateFreshCreate(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "slotstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_creates_file_when_not_exists(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)
        self.assertTrue(os.path.exists(self.xlsx_path))

    def test_headers_correct(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(rows[0][:4], ["##var", "id", "personid", "prefabtype"])
        self.assertIn("face", rows[0])
        face_col = rows[0].index("face") + 1
        self.assertEqual(ws.cell(row=2, column=face_col).value, "int")

    def test_single_row_data(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(rows[4][1], 1000)
        self.assertEqual(rows[4][2], 1)
        self.assertEqual(rows[4][3], 1)
        face_col = rows[0].index("face")
        self.assertEqual(rows[4][face_col], 1)

    def test_multiple_slots(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {
                    "personid": 1,
                    "prefabtype": 1,
                    "slots": {"face": 1, "body": 2, "hair": 0},
                },
            ],
        )
        update_slotstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        face_col = rows[0].index("face")
        body_col = rows[0].index("body")
        hair_col = rows[0].index("hair")
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[4][face_col], 1)
        self.assertEqual(rows[4][body_col], 2)
        self.assertEqual(rows[4][hair_col], 0)

    def test_first_column_none(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertIsNone(rows[4][0])


class TestSlotStateUpdate(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "slotstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_update_existing_slot_value(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": -1}},
            ],
            "update.json",
        )
        update_slotstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        face_col = rows[0].index("face")
        self.assertEqual(rows[4][face_col], -1)
        self.assertEqual(len(rows), 5)

    def test_update_preserves_id(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
                {"personid": 2, "prefabtype": 1, "slots": {"face": 2}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": -1}},
                {"personid": 2, "prefabtype": 1, "slots": {"face": 2}},
            ],
            "update.json",
        )
        update_slotstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(rows[4][1], 1000)
        self.assertEqual(rows[5][1], 2000)

    def test_add_new_person_creates_new_row(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
                {"personid": 2, "prefabtype": 1, "slots": {"face": 2, "hair": 1}},
            ],
            "add.json",
        )
        update_slotstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        face_col = rows[0].index("face")
        self.assertEqual(len(rows), 6)
        self.assertEqual(rows[4][face_col], 1)
        self.assertEqual(rows[5][face_col], 2)

    def test_id_auto_increments_from_max(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
                {"personid": 2, "prefabtype": 1, "slots": {"face": 2}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
                {"personid": 2, "prefabtype": 1, "slots": {"face": 2}},
                {"personid": 3, "prefabtype": 1, "slots": {"face": 1}},
            ],
            "new.json",
        )
        update_slotstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        self.assertEqual(rows[6][1], 3000)


class TestSlotStateCleanup(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "slotstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_no_delete_different_person(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
                {"personid": 2, "prefabtype": 1, "slots": {"face": 2}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 0}},
                {"personid": 2, "prefabtype": 1, "slots": {"face": 0}},
            ],
            "both.json",
        )
        update_slotstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        data_rows = [r for r in rows[4:] if r[1] is not None]
        self.assertEqual(len(data_rows), 2)

    def test_no_delete_different_prefabtype(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
                {"personid": 1, "prefabtype": 2, "slots": {"face": 2}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 0}},
                {"personid": 1, "prefabtype": 2, "slots": {"face": 2}},
            ],
            "one.json",
        )
        update_slotstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        data_rows = [r for r in rows[4:] if r[1] is not None]
        self.assertEqual(len(data_rows), 2)


class TestSlotStateCrossPersonIsolation(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "slotstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_person1_update_does_not_affect_person2(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
                {"personid": 2, "prefabtype": 1, "slots": {"face": 2}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": -1}},
                {"personid": 2, "prefabtype": 1, "slots": {"face": 2}},
            ],
            "update.json",
        )
        update_slotstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        face_col = rows[0].index("face")
        self.assertEqual(rows[4][2], 1)
        self.assertEqual(rows[4][face_col], -1)
        self.assertEqual(rows[5][2], 2)
        self.assertEqual(rows[5][face_col], 2)


class TestSlotStateBOM(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "slotstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_utf8_bom_json_handled(self):
        json_path = _write_json_bom(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        face_col = rows[0].index("face")
        self.assertEqual(rows[4][face_col], 1)


class TestAnimationStateSort(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "animationstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_sort_by_personid_then_prefabtype(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 3, "statename": "idle", "prefabtype": 2},
                {"personid": 1, "statename": "talk", "prefabtype": 2},
                {"personid": 1, "statename": "idle", "prefabtype": 1},
                {"personid": 2, "statename": "idle", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        data_rows = [r for r in rows[4:] if r[1] is not None]
        pids = [r[2] for r in data_rows]
        pts = [r[4] for r in data_rows]
        self.assertEqual(pids, [1, 1, 2, 3])
        self.assertEqual(pts, [1, 2, 1, 2])

    def test_sort_after_update(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 2, "statename": "idle", "prefabtype": 1},
            ],
        )
        update_animationstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "statename": "talk", "prefabtype": 1},
            ],
            "add.json",
        )
        update_animationstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        data_rows = [r for r in rows[4:] if r[1] is not None]
        pids = [r[2] for r in data_rows]
        self.assertEqual(pids, [1, 2])


class TestSlotStateSort(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.xlsx_path = os.path.join(self.temp_dir, "slotstate.xlsx")

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_sort_by_personid_then_prefabtype(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 3, "prefabtype": 1, "slots": {"face": 0}},
                {"personid": 1, "prefabtype": 2, "slots": {"face": 0}},
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
                {"personid": 2, "prefabtype": 1, "slots": {"face": 2}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)
        rows, ws = _read_xlsx(self.xlsx_path)

        data_rows = [r for r in rows[4:] if r[1] is not None]
        pids = [r[2] for r in data_rows]
        pts = [r[3] for r in data_rows]
        self.assertEqual(pids, [1, 1, 2, 3])
        self.assertEqual(pts, [1, 2, 1, 1])

    def test_sort_after_update(self):
        json_path = _write_json(
            self.temp_dir,
            [
                {"personid": 2, "prefabtype": 1, "slots": {"face": 2}},
            ],
        )
        update_slotstate(self.xlsx_path, json_path)

        json_path2 = _write_json(
            self.temp_dir,
            [
                {"personid": 1, "prefabtype": 1, "slots": {"face": 1}},
            ],
            "add.json",
        )
        update_slotstate(self.xlsx_path, json_path2)
        rows, ws = _read_xlsx(self.xlsx_path)

        data_rows = [r for r in rows[4:] if r[1] is not None]
        pids = [r[2] for r in data_rows]
        self.assertEqual(pids, [1, 2])


if __name__ == "__main__":
    unittest.main()
