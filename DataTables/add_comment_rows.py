"""
在 Luban Excel 中补充备注说明行：
1. 从原始 Excel（Assets/Configs/Csv/）读取第1行（备注说明）
2. 在 Luban Excel 的 ##type 行后插入 ##（备注说明）
3. 将原有 ## 注释行改为 ##comment
"""

import openpyxl
import os
import sys
import copy

sys.stdout.reconfigure(encoding="utf-8")

ORIG_DIR = os.path.join("..", "Assets", "Configs", "Csv")
LUBAN_DIR = os.path.join(".", "Datas")

ORIG_TO_LUBAN = {
    ("bag.xlsx", "bag"): ("bag.xlsx", None),
    ("buff.xlsx", "buff"): ("buff.xlsx", None),
    ("condition.xlsx", "condition"): ("condition.xlsx", None),
    ("dialog.xlsx", "dialog"): ("dialog.xlsx", None),
    ("dialog.xlsx", "dialogcontent"): ("dialogcontent.xlsx", None),
    ("item.xlsx", "item"): ("item.xlsx", None),
    ("character.xlsx", "person"): ("person.xlsx", None),
    ("quest.xlsx", "quest"): ("quest.xlsx", None),
    ("quest.xlsx", "questcontext"): ("questcontext.xlsx", None),
}


def read_orig_row1(orig_file, sheet_name):
    path = os.path.join(ORIG_DIR, orig_file)
    if not os.path.exists(path):
        return None
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    row1 = []
    for cell in ws[1]:
        row1.append(cell.value if cell.value is not None else "")
    wb.close()
    non_empty = [x for x in row1 if str(x).strip()]
    if not non_empty:
        return None
    return row1


def process_luban(luban_file, orig_row1):
    path = os.path.join(LUBAN_DIR, luban_file)
    if not os.path.exists(path):
        print(f"  跳过: {luban_file} 不存在")
        return

    wb = openpyxl.load_workbook(path)
    ws = wb.active

    type_row = None
    for r in range(1, 10):
        val = ws.cell(row=r, column=1).value
        if val and str(val).strip() == "##type":
            type_row = r
            break

    if type_row is None:
        print(f"  跳过: 未找到 ##type")
        wb.close()
        return

    ws.insert_rows(type_row + 1)

    for col_idx, val in enumerate(orig_row1, start=1):
        ws.cell(row=type_row + 1, column=col_idx, value=val)
    ws.cell(row=type_row + 1, column=1).value = "##"

    comment_row = type_row + 2
    cell = ws.cell(row=comment_row, column=1)
    if cell.value and str(cell.value).strip() == "##":
        cell.value = "##comment"

    wb.save(path)
    print(f"  已处理: 插入 ## 备注行，原有注释改为 ##comment")
    wb.close()


def main():
    for (orig_file, orig_sheet), (luban_file, _) in ORIG_TO_LUBAN.items():
        print(f"\n{orig_file}[{orig_sheet}] -> {luban_file}")
        row1 = read_orig_row1(orig_file, orig_sheet)
        if row1 is None:
            print(f"  跳过: 原始 Row1 为空或文件不存在")
            continue
        preview = [str(x)[:20] for x in row1[:5] if str(x).strip()]
        print(f"  备注内容: {preview}")
        process_luban(luban_file, row1)

    print("\n=== 完成 ===")


if __name__ == "__main__":
    main()
