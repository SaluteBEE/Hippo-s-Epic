import sys
import json
import openpyxl
import os
import tempfile


def safe_save(wb, xlsx_path):
    tmp = xlsx_path + ".tmp"
    wb.save(tmp)
    if os.path.exists(xlsx_path):
        try:
            os.remove(xlsx_path)
        except PermissionError:
            pass
    try:
        os.replace(tmp, xlsx_path)
    except PermissionError:
        print(f"WARN: Cannot replace {xlsx_path}, file locked by Unity")
        print(f"DATA saved to {tmp}, close Unity or re-import to apply")


def _sort_data_rows(ws, key_cols):
    data_start = 5
    data_end = ws.max_row
    if data_end < data_start:
        return
    rows_data = []
    for row_idx in range(data_start, data_end + 1):
        row_cells = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]
        sort_key = tuple(row_cells[k - 1] if row_cells[k - 1] is not None else 0 for k in key_cols)
        rows_data.append((sort_key, row_cells))
    rows_data.sort(key=lambda x: x[0])
    for i, (_, row_cells) in enumerate(rows_data):
        for col_idx, val in enumerate(row_cells, 1):
            ws.cell(row=data_start + i, column=col_idx).value = val


def _delete_rows(ws, row_indices):
    if not row_indices:
        return 0
    for row_idx in sorted(row_indices, reverse=True):
        ws.delete_rows(row_idx)
    return len(row_indices)


def update_animationstate(xlsx_path, json_path):
    with open(json_path, "r", encoding="utf-8-sig") as f:
        rows = json.load(f)

    valid_keys = {(row["personid"], row["statename"], row.get("prefabtype", 1)) for row in rows}
    scope_combos = {(row["personid"], row.get("prefabtype", 1)) for row in rows}

    if os.path.exists(xlsx_path):
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active

        comp_col = None
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=1, column=c).value == "compositionname":
                comp_col = c
                break
        if comp_col is not None:
            ws.delete_cols(comp_col)

        existing = {}
        for row_idx in range(5, ws.max_row + 1):
            pid = ws.cell(row=row_idx, column=3).value
            sn = ws.cell(row=row_idx, column=4).value
            pt = ws.cell(row=row_idx, column=5).value
            if pid is not None and sn is not None and pt is not None:
                existing[(int(pid), str(sn), int(pt))] = row_idx

        stale_rows = []
        for key, row_idx in existing.items():
            if (key[0], key[2]) in scope_combos:
                if key not in valid_keys:
                    stale_rows.append(row_idx)
        deleted = _delete_rows(ws, stale_rows)

        existing = {}
        for row_idx in range(5, ws.max_row + 1):
            pid = ws.cell(row=row_idx, column=3).value
            sn = ws.cell(row=row_idx, column=4).value
            pt = ws.cell(row=row_idx, column=5).value
            if pid is not None and sn is not None and pt is not None:
                existing[(int(pid), str(sn), int(pt))] = row_idx
        max_id_by_person = {}
        for r in range(5, ws.max_row + 1):
            v = ws.cell(row=r, column=2).value
            pid_v = ws.cell(row=r, column=3).value
            if v is not None and pid_v is not None:
                pid = int(pid_v)
                offset = int(v) - pid * 1000
                cur = max_id_by_person.get(pid, 0)
                if offset >= cur:
                    max_id_by_person[pid] = offset + 1
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "animationstate"
        ws.append(["##var", "id", "personid", "statename", "prefabtype", "slotstateid"])
        ws.append(["##type", "int", "int#ref=cfg.person.TbPerson", "string", "int", "int"])
        ws.append(["##", None, None, None, None, None])
        ws.append(["##comment", "动画状态ID", "角色外键", "状态名(=组合名)", "预制体类型(1=对话,2=战斗)", "关联插槽状态ID(0=不变)"])
        existing = {}
        max_id_by_person = {}
        deleted = 0

    updated = 0
    added = 0
    for row in rows:
        pid = row["personid"]
        key = (pid, row["statename"], row.get("prefabtype", 1))
        if key in existing:
            row_idx = existing[key]
            ws.cell(row=row_idx, column=6).value = row.get("slotstateid", 0)
            updated += 1
        else:
            next_id = pid * 1000 + max_id_by_person.get(pid, 0)
            ws.append(
                [None, next_id, pid, row["statename"], row.get("prefabtype", 1), row.get("slotstateid", 0)]
            )
            existing[key] = ws.max_row
            max_id_by_person[pid] = max_id_by_person.get(pid, 0) + 1
            added += 1

    _sort_data_rows(ws, key_cols=[3, 5])

    safe_save(wb, xlsx_path)
    parts = [f"updated={updated}", f"added={added}"]
    if deleted > 0:
        parts.append(f"deleted={deleted}")
    print(f"OK animationstate: {', '.join(parts)}")


def _get_slot_columns(ws):
    result = []
    col = 5
    while True:
        val = ws.cell(row=1, column=col).value
        if val is None or str(val).startswith("##"):
            break
        result.append((col, str(val)))
        col += 1
    return result


def _ensure_slot_columns(ws, slot_names):
    existing = {name: col for col, name in _get_slot_columns(ws)}
    max_col = max((col for col, _ in _get_slot_columns(ws)), default=4) + 1 if _get_slot_columns(ws) else 5
    for name in slot_names:
        if name not in existing:
            col = max_col
            ws.cell(row=1, column=col, value=name)
            ws.cell(row=2, column=col, value="int")
            ws.cell(row=3, column=col, value=None)
            ws.cell(row=4, column=col, value=f"{name}插槽")
            existing[name] = col
            max_col += 1
    return existing


def update_slotstate(xlsx_path, json_path):
    with open(json_path, "r", encoding="utf-8-sig") as f:
        rows = json.load(f)

    all_slot_names = []
    for row in rows:
        for k in row.get("slots", {}):
            if k not in all_slot_names:
                all_slot_names.append(k)

    valid_combos = {(row["personid"], row.get("prefabtype", 1)) for row in rows}

    if os.path.exists(xlsx_path):
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        slot_col_map = _ensure_slot_columns(ws, all_slot_names)
        existing = {}
        for row_idx in range(5, ws.max_row + 1):
            rid = ws.cell(row=row_idx, column=2).value
            if rid is not None and isinstance(rid, (int, float)):
                existing[int(rid)] = row_idx

        stale_rows = []
        deleted = 0

        existing = {}
        for row_idx in range(5, ws.max_row + 1):
            rid = ws.cell(row=row_idx, column=2).value
            if rid is not None and isinstance(rid, (int, float)):
                existing[int(rid)] = row_idx
        max_id_by_person = {}
        for rid, row_idx_existing in existing.items():
            pid_v = ws.cell(row=row_idx_existing, column=3).value
            if pid_v is not None:
                pid = int(pid_v)
                offset = rid - pid * 1000
                cur = max_id_by_person.get(pid, 0)
                if offset >= cur:
                    max_id_by_person[pid] = offset + 1
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "slotstate"
        header_var = ["##var", "id", "personid", "prefabtype"] + all_slot_names
        header_type = ["##type", "int", "int#ref=cfg.person.TbPerson", "int"] + [
            "int"
        ] * len(all_slot_names)
        header_comment = ["##comment", "状态ID", "角色外键", "预制体类型(1=对话,2=战斗)"] + [
            "-1=隐藏,0=不变,N=附件N" for _ in all_slot_names
        ]
        ws.append(header_var)
        ws.append(header_type)
        ws.append(["##"] + [None] * (3 + len(all_slot_names)))
        ws.append(header_comment)
        slot_col_map = {name: 5 + i for i, name in enumerate(all_slot_names)}
        existing = {}
        max_id_by_person = {}
        deleted = 0

    updated = 0
    added = 0
    for row in rows:
        pid = row["personid"]
        pt = row.get("prefabtype", 1)
        slots = row.get("slots", {})

        matched_id = None
        for rid, row_idx in existing.items():
            exist_pid = ws.cell(row=row_idx, column=3).value
            exist_pt = ws.cell(row=row_idx, column=4).value
            if exist_pid == pid and exist_pt == pt:
                matched_id = rid
                break

        if matched_id is not None:
            row_idx = existing[matched_id]
            for slot_name, val in slots.items():
                if slot_name in slot_col_map:
                    ws.cell(row=row_idx, column=slot_col_map[slot_name]).value = int(val)
            updated += 1
        else:
            next_id = pid * 1000 + max_id_by_person.get(pid, 0)
            data = [None, next_id, pid, pt]
            data.extend([0] * len(slot_col_map))
            ws.append(data)
            target_row = ws.max_row
            for slot_name, val in slots.items():
                if slot_name in slot_col_map:
                    ws.cell(row=target_row, column=slot_col_map[slot_name]).value = int(val)
            existing[next_id] = target_row
            max_id_by_person[pid] = max_id_by_person.get(pid, 0) + 1
            added += 1

    _sort_data_rows(ws, key_cols=[3, 4])

    safe_save(wb, xlsx_path)
    parts = [f"updated={updated}", f"added={added}"]
    if deleted > 0:
        parts.append(f"deleted={deleted}")
    print(f"OK slotstate: {', '.join(parts)}")


if __name__ == "__main__":
    mode = sys.argv[1]
    xlsx_path = sys.argv[2]
    json_path = sys.argv[3]

    if mode == "animationstate":
        update_animationstate(xlsx_path, json_path)
    elif mode == "slotstate":
        update_slotstate(xlsx_path, json_path)
    else:
        print(f"ERR unknown mode: {mode}")
        sys.exit(1)
